"""Read-only report/export surface for the payment database.

This module deliberately does not expose generic writes. Financial tables remain
owned by the payment domain APIs and ledger invariants.
"""
import io
from datetime import datetime

from app.api.admin import require_admin
from app.core.db import Base, SessionLocal
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from sqlalchemy import String, and_, func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

router = APIRouter(prefix="/data-reports", tags=["admin-data-reports"])
MAX_ROWS = 10000

async def db():
    async with SessionLocal() as s:
        yield s

def _admin(user):
    if "platform_admin" not in set(user.get("roles", [])):
        raise HTTPException(403, "platform_admin required")

def _table(name):
    table = Base.metadata.tables.get(name)
    if table is None:
        raise HTTPException(404, f"No table named '{name}'")
    return table

def _row(row, table):
    return [str(getattr(row, c.name)) if getattr(row, c.name) is not None else "" for c in table.columns]

async def _rows(table, search, s):
    predicate = None
    if search:
        predicate = or_(*[c.cast(String).ilike(f"%{search}%") for c in table.columns])
    stmt = select(table)
    if predicate is not None:
        stmt = stmt.where(predicate)
    return table, (await s.execute(stmt.limit(MAX_ROWS))).fetchall()

@router.get("/{table_name}/xlsx")
async def export_xlsx(table_name: str, search: str | None = Query(None, max_length=200), user=Depends(require_admin), s: AsyncSession = Depends(db)):
    _admin(user)
    table, rows = await _rows(_table(table_name), search, s)
    wb = Workbook()
    ws = wb.active
    ws.title = table_name[:31]
    ws.append([c.name for c in table.columns])
    for row in rows:
        ws.append(_row(row, table))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename={table_name}.xlsx"})

@router.get("/{table_name}/pdf")
async def export_pdf(table_name: str, search: str | None = Query(None, max_length=200), user=Depends(require_admin), s: AsyncSession = Depends(db)):
    _admin(user)
    table, rows = await _rows(_table(table_name), search, s)
    headers = [c.name for c in table.columns]
    data = [headers] + [_row(r, table) for r in rows]
    out = io.BytesIO()
    doc = SimpleDocTemplate(out, pagesize=landscape(letter), title=f"Shopnoltd {table_name} report")
    rendered = Table(data, repeatRows=1)
    rendered.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.grey),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 6),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
    ]))
    doc.build([rendered])
    out.seek(0)
    return StreamingResponse(out, media_type="application/pdf", headers={"Content-Disposition": f"attachment; filename={table_name}.pdf"})

@router.get("/summary")
async def summary(user=Depends(require_admin), s: AsyncSession = Depends(db)):
    _admin(user)
    result = {}
    for name, table in sorted(Base.metadata.tables.items()):
        try:
            result[name] = int((await s.execute(select(func.count()).select_from(table))).scalar_one())
        except Exception:
            result[name] = None
    return {"generated_at": datetime.utcnow().isoformat(), "tables": result}
